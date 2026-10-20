"""Excel import: upload, the worker job, and what each row turns into."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from config import settings
from database.config import SessionLocal, bind_tenant
from models import Book, BookSource, BookStatus, Job, MemberRole, Tenant, User
from tests.factories import add_member, make_book, make_tenant, make_user, sign_in
from workers import main as worker

pytestmark = pytest.mark.db

HEADER = ["ISBN", "Kitob nomi", "Muallif", "Nashriyot", "Narx (so'm)", "Soni"]


def xlsx(*rows: list[object], header: list[str] = HEADER) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def upload(data: bytes) -> dict:
    return {
        "file": ("kitoblar.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }


async def store(db: AsyncSession) -> tuple[Tenant, User]:
    owner = await make_user(db, "Owner")
    return await make_tenant(db, "varaq", owner), owner


async def run_import(client: AsyncClient, tenant: Tenant, data: bytes, mode: str | None = None) -> dict:
    params = {"quantity_mode": mode} if mode else None
    created = await client.post(f"/tenants/{tenant.id}/inventory/imports", files=upload(data), params=params)
    assert created.status_code == 202, created.text
    assert created.json()["status"] == "queued"
    while await worker.process_one():
        pass
    done = await client.get(f"/tenants/{tenant.id}/inventory/imports/{created.json()['id']}")
    return done.json()


async def shelf(client: AsyncClient, tenant: Tenant) -> dict[str, tuple[int, int]]:
    items = (await client.get(f"/tenants/{tenant.id}/inventory", params={"size": 100})).json()["items"]
    return {row["book"]["title"]: (row["price"], row["quantity"]) for row in items}


async def test_template_has_the_columns(client: AsyncClient, db: AsyncSession) -> None:
    tenant, owner = await store(db)
    await sign_in(client, owner)
    resp = await client.get(f"/tenants/{tenant.id}/inventory/import-template")
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(resp.content)).worksheets[0]
    assert [cell.value for cell in sheet[1]] == [
        "isbn", "title", "authors", "publisher", "language", "script", "price", "quantity",
    ]  # fmt: skip


async def test_import_matches_creates_and_reports(client: AsyncClient, db: AsyncSession) -> None:
    tenant, owner = await store(db)
    await make_book(db, "O'tkan kunlar", ["Abdulla Qodiriy"], isbn="9780306406157")
    await make_book(db, "Sariq devni minib", ["Xudoyberdi To'xtaboyev"], publisher="Yangi asr")
    await sign_in(client, owner)

    data = xlsx(
        [9780306406157, None, None, None, "45 000", 3],  # by ISBN (typed as a number)
        [None, "Sariq devni minib", "Xudoyberdi Toʻxtaboyev", "Yangi asr", 30000, 2],  # by edition
        [None, "Yangi kitob", "Muallif Bir; Muallif Ikki", None, "12000.50", None],  # new -> pending
        [None, "Narxsiz", "Kimdir", None, "bepul", 1],  # bad price
        ["9781234567897", None, None, None, 1000, 1],  # unknown ISBN, no title
        ["123", "Yomon ISBN", "Kimdir", None, 1000, 1],  # invalid ISBN
        [None, "Manfiy", "Kimdir", None, 1000, -2],  # bad quantity
    )
    report = await run_import(client, tenant, data)
    assert report["status"] == "done"
    assert (report["total_rows"], report["processed_rows"]) == (7, 7)
    r = report["report"]
    assert (r["books_matched"], r["books_created"], r["items_added"], r["failed_rows"]) == (2, 1, 3, 4)
    assert [(e["line"], e["code"]) for e in r["errors"]] == [
        (5, "import_row_bad_price"),
        (6, "import_row_book_unknown"),
        (7, "isbn_invalid"),
        (8, "import_row_bad_quantity"),
    ]
    assert all(e["message"] for e in r["errors"])  # human text, in the uploader's language

    assert await shelf(client, tenant) == {
        "O'tkan kunlar": (4_500_000, 3),
        "Sariq devni minib": (3_000_000, 2),
        "Yangi kitob": (1_200_050, 0),
    }
    new_book = await db.scalar(select(Book).where(Book.title == "Yangi kitob"))
    assert (new_book.status, new_book.submitted_via) == (BookStatus.PENDING, BookSource.IMPORT)

    # The same file again ("add", the default): the quantities arrive once more, books are not duplicated.
    again = await run_import(client, tenant, data)
    assert again["quantity_mode"] == "add"
    assert again["report"]["items_updated"] == 3 and again["report"]["items_added"] == 0
    assert await db.scalar(select(func.count()).select_from(Book).where(Book.title == "Yangi kitob")) == 1
    assert (await shelf(client, tenant))["O'tkan kunlar"] == (4_500_000, 6)

    # A stocktake ("set"): the file says how many there are now.
    counted = await run_import(client, tenant, data, mode="set")
    assert counted["quantity_mode"] == "set"
    assert (await shelf(client, tenant))["O'tkan kunlar"] == (4_500_000, 3)


async def test_rejected_book_is_not_submitted_again(client: AsyncClient, db: AsyncSession) -> None:
    tenant, owner = await store(db)
    admin = await make_user(db, "Admin", is_platform_admin=True)
    await sign_in(client, owner)
    book = (
        await client.post(
            f"/tenants/{tenant.id}/catalog/books", json={"title": "Takror", "authors": ["Kimdir Bir"]}
        )
    ).json()
    await sign_in(client, admin)
    await client.post(f"/system/catalog/books/{book['id']}/reject", json={"reason": "Qalbaki nashr"})

    await sign_in(client, owner)
    report = await run_import(client, tenant, xlsx([None, "Takror", "Kimdir Bir", None, 1000, 1]))
    [error] = report["report"]["errors"]
    assert error["code"] == "import_row_book_rejected"
    assert "Qalbaki nashr" in error["message"]


async def test_import_pending_limit_is_separate_from_manual(
    client: AsyncClient, db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "MAX_PENDING_BOOKS_FROM_IMPORT", 1)
    monkeypatch.setattr(settings, "MAX_PENDING_BOOKS_PER_STORE", 1)
    tenant, owner = await store(db)
    await sign_in(client, owner)
    report = await run_import(
        client, tenant, xlsx([None, "Bir", "Muallif A", None, 1, 1], [None, "Ikki", "Muallif B", None, 1, 1])
    )
    assert report["report"]["books_created"] == 1
    assert report["report"]["errors"][0]["code"] == "pending_books_limit"
    # Adding one by hand still works: the import did not use up the manual quota.
    manual = await client.post(
        f"/tenants/{tenant.id}/catalog/books", json={"title": "Uch", "authors": ["Muallif C"]}
    )
    assert manual.status_code == 201


@pytest.mark.parametrize(
    ("data", "reason"),
    [
        (b"not a spreadsheet", "not_xlsx"),
        (xlsx(header=["foo", "bar"]), "no_header"),
        (xlsx(header=["isbn", "title"]), "missing_columns"),
        (xlsx(header=HEADER), "empty"),
    ],
)
async def test_unreadable_files_are_refused_at_upload(
    client: AsyncClient, db: AsyncSession, data: bytes, reason: str
) -> None:
    tenant, owner = await store(db)
    await sign_in(client, owner)
    resp = await client.post(f"/tenants/{tenant.id}/inventory/imports", files=upload(data))
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "import_file_invalid"
    assert resp.json()["detail"]["reason"] == reason
    assert await db.scalar(select(func.count()).select_from(Job)) == 0


async def test_too_many_rows(client: AsyncClient, db: AsyncSession, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "MAX_IMPORT_ROWS", 2)
    tenant, owner = await store(db)
    await sign_in(client, owner)
    data = xlsx(*[[None, f"Kitob {i}", "Kimdir", None, 1, 1] for i in range(3)])
    resp = await client.post(f"/tenants/{tenant.id}/inventory/imports", files=upload(data))
    assert resp.json()["detail"]["reason"] == "too_many_rows"


async def test_imports_are_the_owners(client: AsyncClient, db: AsyncSession) -> None:
    tenant, _ = await store(db)
    staff = await make_user(db, "Staff")
    await add_member(db, tenant, staff, MemberRole.STAFF)
    await sign_in(client, staff)
    resp = await client.post(
        f"/tenants/{tenant.id}/inventory/imports", files=upload(xlsx([None, "A", "B", None, 1, 1]))
    )
    assert resp.status_code == 403


async def test_add_mode_adds_to_what_is_on_the_shelf(client: AsyncClient, db: AsyncSession) -> None:
    """The reported bug: 5 on the shelf, the file says "2 arrived" -> 7, not 2."""
    tenant, owner = await store(db)
    await make_book(db, "O'tkan kunlar", ["Abdulla Qodiriy"], isbn="9780306406157")
    await sign_in(client, owner)
    await run_import(client, tenant, xlsx([9780306406157, None, None, None, 45000, 5]))
    await run_import(client, tenant, xlsx([9780306406157, None, None, None, 47000, 2]))
    assert (await shelf(client, tenant))["O'tkan kunlar"] == (4_700_000, 7)
    bad_mode = await client.post(
        f"/tenants/{tenant.id}/inventory/imports",
        files=upload(xlsx([9780306406157, None, None, None, 1, 1])),
        params={"quantity_mode": "double"},
    )
    assert bad_mode.status_code == 422


async def test_readable_row_problems(client: AsyncClient, db: AsyncSession) -> None:
    tenant, owner = await store(db)
    await sign_in(client, owner)
    header = ["isbn", "title", "authors", "language", "script", "price", "quantity"]
    report = await run_import(
        client,
        tenant,
        xlsx(
            ["978-9943-55-123-4", "Kitob A", "Muallif A", None, None, 1000, 1],  # wrong check digit
            [None, "Kitob B", "Muallif B", "O'zbek", "Lotin", 1000, 1],  # words, not codes: fine
            [None, "Kitob C", "Muallif C", "Русский", "kirill", 1000, 1],  # fine
            [None, "Kitob D", "Muallif D", "klingon", None, 1000, 1],  # unknown language
            [None, "Kitob E", "Muallif E", None, None, 0, 1],  # price 0
            header=header,
        ),
    )
    errors = {e["line"]: e for e in report["report"]["errors"]}
    assert errors[2]["code"] == "isbn_check_digit"
    assert errors[2]["message"].startswith("«978-9943-55-123-4»:") and " 7 " in errors[2]["message"]
    assert errors[5]["code"] == "book_language_invalid"
    assert errors[6]["code"] == "import_row_bad_price"
    assert set(errors) == {2, 5, 6}
    books = {b.title: (b.language, b.script) for b in await db.scalars(select(Book))}
    assert books["Kitob B"] == ("uz", "latn") and books["Kitob C"] == ("ru", "cyrl")


async def test_a_retried_import_continues_and_adds_nothing_twice(
    client: AsyncClient, db: AsyncSession
) -> None:
    """A worker that died after committing row 1 must not add row 1 again."""
    from models import InventoryImport

    tenant, owner = await store(db)
    await make_book(db, "Bir", ["Muallif A"], isbn="9780306406157")
    await make_book(db, "Ikki", ["Muallif B"])
    await sign_in(client, owner)
    data = xlsx([9780306406157, None, None, None, 1000, 4], [None, "Ikki", "Muallif B", None, 1000, 2])
    created = (await client.post(f"/tenants/{tenant.id}/inventory/imports", files=upload(data))).json()
    # As if an earlier attempt had committed row 1 (shelf + progress) and then crashed.
    await make_item_direct(tenant, "Bir", quantity=4)
    await bind_tenant(db, tenant.id)  # inventory_imports is a store table (RLS)
    await db.execute(
        update(InventoryImport)
        .where(InventoryImport.id == created["id"])
        .values(processed_rows=1, report={"books_matched": 1, "items_added": 1, "errors": []})
    )
    await db.commit()
    while await worker.process_one():
        pass
    done = (await client.get(f"/tenants/{tenant.id}/inventory/imports/{created['id']}")).json()
    assert done["report"]["books_matched"] == 2 and done["report"]["items_added"] == 2
    assert await shelf(client, tenant) == {"Bir": (100_000, 4), "Ikki": (100_000, 2)}


async def make_item_direct(tenant: Tenant, title: str, quantity: int) -> None:
    from models import InventoryItem

    async with SessionLocal() as session:
        await bind_tenant(session, tenant.id)
        book_id = await session.scalar(select(Book.id).where(Book.title == title))
        session.add(InventoryItem(tenant_id=tenant.id, book_id=book_id, price=100_000, quantity=quantity))
        await session.commit()
