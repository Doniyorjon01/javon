"""
Excel (.xlsx) template and reader for inventory imports
=======================================================
The store downloads the template, fills one row per book and uploads it.
Columns (first row; Uzbek, Russian or English headers are all recognised):

    isbn       ISBN-10 or ISBN-13; may be empty
    title      required when the ISBN is not in the catalog yet
    authors    required with title; several separated by ";" (or ",")
    publisher, language (uz|ru|en|other, or "o'zbek", "rus"...), script (latn|cyrl, or "lotin", "kirill")
    price      in SO'M, as people write it: 45000, "45 000", 45000.00   (required)
    quantity   how many; added to the shelf or replacing it (InventoryImport.quantity_mode)

Parsing only: what a row MEANS (match a catalog book, create one, put it on the
shelf) is workers/inventory_import.py.

Security: openpyxl reads XML; defusedxml (in requirements.txt) makes it refuse
entity-expansion attacks. read_only mode streams rows instead of loading the
whole sheet, and MAX_IMPORT_ROWS bounds the work. CPU-bound: run in a thread.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException

from utils.translit import search_key

COLUMNS = ("isbn", "title", "authors", "publisher", "language", "script", "price", "quantity")

# Header text (normalized with search_key) -> column.
_ALIASES: dict[str, tuple[str, ...]] = {
    "isbn": ("isbn", "shtrix kod", "shtrixkod", "barcode", "штрих код"),
    "title": ("title", "nomi", "kitob nomi", "nom", "название", "наименование"),
    "authors": ("authors", "author", "muallif", "mualliflar", "автор", "авторы"),
    "publisher": ("publisher", "nashriyot", "издательство"),
    "language": ("language", "til", "язык"),
    "script": ("script", "yozuv", "алфавит"),
    "price": ("price", "narx", "narxi", "narx som", "narxi som", "цена", "цена сум"),
    "quantity": ("quantity", "soni", "qoldiq", "miqdor", "количество", "остаток"),
}
HEADER_TO_COLUMN = {search_key(alias): column for column, aliases in _ALIASES.items() for alias in aliases}

MAX_CELL_LENGTH = 5000


class ExcelProblem(StrEnum):
    NOT_XLSX = "not_xlsx"
    NO_HEADER = "no_header"
    MISSING_COLUMNS = "missing_columns"  # needs price and (isbn or title)
    TOO_MANY_ROWS = "too_many_rows"
    EMPTY = "empty"


class ExcelError(ValueError):
    def __init__(self, problem: ExcelProblem) -> None:
        super().__init__(problem.value)
        self.problem = problem


@dataclass(frozen=True, slots=True)
class Row:
    line: int  # 1-based line number in the sheet, as the store sees it
    values: dict[str, str]  # column -> trimmed text; empty cells left out


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 9780306406157.0 -> "9780306406157" (ISBN typed as a number)
    return str(value).strip()[:MAX_CELL_LENGTH]


def read_rows(data: bytes, max_rows: int) -> list[Row]:
    """All non-empty data rows of the first sheet. Raises ExcelError."""
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExcelError(ExcelProblem.NOT_XLSX)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise ExcelError(ExcelProblem.NOT_XLSX) from None

    try:
        sheet = workbook.worksheets[0]
        rows = enumerate(sheet.iter_rows(values_only=True), start=1)
        header: dict[int, str] = {}
        for _, header_cells in rows:  # the first non-empty row is the header
            texts = [_cell_text(cell) for cell in header_cells]
            if any(texts):
                header = {
                    index: HEADER_TO_COLUMN[key]
                    for index, text in enumerate(texts)
                    if (key := search_key(text)) in HEADER_TO_COLUMN
                }
                break
        if not header:
            raise ExcelError(ExcelProblem.NO_HEADER)
        found = set(header.values())
        if "price" not in found or not ({"isbn", "title"} & found):
            raise ExcelError(ExcelProblem.MISSING_COLUMNS)

        result: list[Row] = []
        for line, cells in rows:
            values = {
                column: text
                for index, column in header.items()
                if index < len(cells) and (text := _cell_text(cells[index]))
            }
            if not values:
                continue
            if len(result) >= max_rows:
                raise ExcelError(ExcelProblem.TOO_MANY_ROWS)
            result.append(Row(line=line, values=values))
    finally:
        workbook.close()

    if not result:
        raise ExcelError(ExcelProblem.EMPTY)
    return result


_MONEY_JUNK = re.compile(
    r"[\s\u00a0\u202f'\u2019`]|so['\u2019`\u02bb]?m|\u0441\u0443\u043c|sum|uzs", re.IGNORECASE
)


def parse_som(text: str) -> int | None:
    """ "45 000", "45000", "45,000", "45000.50", "45 000 so'm" -> tiyin. None if not a sum."""
    cleaned = _MONEY_JUNK.sub("", text)
    if re.fullmatch(r"\d{1,3}(,\d{3})+", cleaned):  # 45,000 = thousands separator
        cleaned = cleaned.replace(",", "")
    cleaned = cleaned.replace(",", ".")
    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    if not amount.is_finite() or amount < 0:
        return None
    return int((amount * 100).quantize(Decimal(1)))


def parse_quantity(text: str) -> int | None:
    cleaned = text.replace(" ", "").replace("\u00a0", "")
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        return None
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        return None
    return int(number)


# What people type in the language / script columns (normalized with search_key,
# so "O'zbek", "o‘zbekcha" and "Ўзбек" all match).
_LANGUAGES = {
    "uz": ("uz", "uzb", "o'zbek", "o'zbekcha", "uzbek", "ўзбек", "ўзбекча", "узбекский", "узб"),
    "ru": ("ru", "rus", "ruscha", "russian", "русский", "рус", "rus tili"),
    "en": ("en", "eng", "ingliz", "inglizcha", "english", "английский", "англ"),
    "other": ("other", "boshqa", "другой", "прочий"),
}
_SCRIPTS = {
    "latn": ("latn", "latin", "lotin", "lotincha", "латиница", "лотин", "lat"),
    "cyrl": ("cyrl", "cyrillic", "kirill", "kirillcha", "кириллица", "кирилл", "kir"),
    "other": ("other", "boshqa", "другой"),
}
_LANGUAGE_BY_KEY = {search_key(alias): code for code, aliases in _LANGUAGES.items() for alias in aliases}
_SCRIPT_BY_KEY = {search_key(alias): code for code, aliases in _SCRIPTS.items() for alias in aliases}


def parse_language(text: str) -> str | None:
    """ "O'zbek", "uz", "русский" -> "uz" / "ru" ...; None if unknown."""
    return _LANGUAGE_BY_KEY.get(search_key(text))


def parse_script(text: str) -> str | None:
    """ "lotin", "Кириллица", "cyrl" -> "latn" / "cyrl"; None if unknown."""
    return _SCRIPT_BY_KEY.get(search_key(text))


def split_authors(text: str) -> list[str]:
    separator = ";" if ";" in text else ","
    return [name.strip() for name in text.split(separator) if name.strip()]


# --------------------------------------------------------------------------- #
# Template
# --------------------------------------------------------------------------- #
_HELP = (
    ("isbn", "Kitob orqasidagi shtrix-kod raqami (978...). Bo'lmasa bo'sh qoldiring."),
    ("title", "Kitob nomi. ISBN katalogda bo'lmasa yoki ISBN yo'q bo'lsa majburiy."),
    ("authors", "Muallif(lar). Bir nechta bo'lsa ; bilan ajrating: Abdulla Qodiriy; Cho'lpon"),
    ("publisher", "Nashriyot (ixtiyoriy)."),
    ("language", "O'zbek (uz), rus (ru), ingliz (en) yoki boshqa. Bo'sh bo'lsa o'zbek."),
    ("script", "Lotin (latn) yoki kirill (cyrl). Bo'sh bo'lsa nomidan aniqlanadi."),
    ("price", "Narx, so'mda: 45000"),
    ("quantity", "Nechta. Yuklashda tanlaysiz: qoldiqqa qo'shiladi yoki qoldiq shu songa almashtiriladi."),
)


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kitoblar"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFGH", (18, 40, 30, 22, 10, 10, 12, 10), strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"

    help_sheet = workbook.create_sheet("Yo'riqnoma")
    help_sheet.append(["Ustun", "Nima yoziladi"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for row in _HELP:
        help_sheet.append(list(row))
    help_sheet.column_dimensions["A"].width = 12
    help_sheet.column_dimensions["B"].width = 80

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
